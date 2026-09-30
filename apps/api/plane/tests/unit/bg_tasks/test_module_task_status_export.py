# Copyright (c) 2023-present Gizmo Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only
# See the LICENSE file for details.

import csv
import io
import zipfile
from datetime import date
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from plane.bgtasks.export_task import issue_export_task
from plane.db.models import ExporterHistory, Issue, Module, ModuleIssue, Project, ProjectMember, State
from plane.utils.porters.exporter import DataExporter
from plane.utils.porters.formatters import XLSXFormatter
from plane.utils.porters.serializers.module_task_status import ModuleTaskStatusExportSerializer


HEADERS = ["Модуль", "Задача", "Статус"]


@pytest.mark.unit
def test_empty_project_report_keeps_column_headers():
    for provider in ("csv", "xlsx"):
        filename, content = DataExporter(
            ModuleTaskStatusExportSerializer,
            format_type=provider,
            headers=HEADERS,
        ).export("report", [])

        assert filename == f"report.{provider}"
        if provider == "csv":
            assert list(csv.reader(io.StringIO(content))) == [HEADERS]
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True)
            assert list(workbook.active.values) == [tuple(HEADERS)]
            workbook.close()


@pytest.mark.unit
def test_spreadsheet_treats_task_titles_as_text():
    content = XLSXFormatter(headers=HEADERS).encode([{"Модуль": "", "Задача": "=1+1", "Статус": ""}])
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    assert workbook.active["B2"].value == "'=1+1"
    assert workbook.active["B2"].data_type == "s"
    workbook.close()


@pytest.mark.unit
@pytest.mark.django_db
def test_project_report_exports_every_issue_once_with_current_modules(workspace, create_user):
    project = Project.objects.create(name="Alpha", identifier="ALPHA", workspace=workspace)
    ProjectMember.objects.create(project=project, member=create_user, role=20)
    state = State.objects.create(project=project, name="В работе", color="#123456")
    first, _ = Issue.objects.bulk_create(
        [
            Issue(project=project, workspace=workspace, state=state, name="Первая", sequence_id=1),
            Issue(
                project=project,
                workspace=workspace,
                state=state,
                name="Без модуля",
                sequence_id=2,
                archived_at=date(2026, 1, 1),
            ),
        ]
    )
    zeta = Module.objects.create(project=project, name="Зета")
    alpha = Module.objects.create(project=project, name="Альфа")
    ModuleIssue.objects.create(project=project, issue=first, module=zeta)
    ModuleIssue.objects.create(project=project, issue=first, module=alpha)

    history = ExporterHistory.objects.create(
        workspace=workspace,
        project=[project.id],
        initiated_by=create_user,
        provider="csv",
        name="module_task_status",
    )
    with patch("plane.bgtasks.export_task.upload_to_s3") as upload:
        issue_export_task(
            provider="csv",
            workspace_id=workspace.id,
            project_ids=[str(project.id)],
            token_id=history.token,
            multiple=False,
            slug=workspace.slug,
            layout="module_task_status",
        )

    upload.assert_called_once()
    with zipfile.ZipFile(upload.call_args.args[0]) as archive:
        assert archive.namelist() == [f"{workspace.slug}-{project.id}-module-task-status.csv"]
        rows = list(csv.DictReader(io.StringIO(archive.read(archive.namelist()[0]).decode("utf-8"))))

    assert rows == [
        {"Модуль": "Альфа, Зета", "Задача": "Первая", "Статус": "В работе"},
        {"Модуль": "", "Задача": "Без модуля", "Статус": "В работе"},
    ]
